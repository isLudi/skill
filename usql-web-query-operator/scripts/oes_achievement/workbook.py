"""Verify native XLSX downloads and hand a stable file manifest to later Base work."""

from __future__ import annotations

from collections import Counter
from decimal import Decimal, InvalidOperation
import hashlib
from io import BytesIO
import json
from pathlib import Path
from typing import Any
import warnings
from copy import copy
import sys

from _shared.errors import UsageError

SKILLS_ROOT = Path(__file__).resolve().parents[3]


def validate_output_dir(path: Path) -> Path:
    resolved = path.expanduser().resolve()
    if resolved == SKILLS_ROOT or SKILLS_ROOT in resolved.parents:
        raise UsageError("Downloads, receipts and query data must stay outside the skills repository.")
    return resolved


def file_info(path: Path) -> dict[str, Any]:
    return {"path": str(path.resolve()), "bytes": path.stat().st_size,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def atomic_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)


def verify_xlsx(data: bytes, expected_count: int, query_rows: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise UsageError("Native XLSX verification requires openpyxl in the configured Python runtime.") from exc
    try:
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Workbook contains no default style, apply openpyxl's default", category=UserWarning)
            book = load_workbook(BytesIO(data), read_only=True, data_only=True)
        try:
            if len(book.worksheets) != 1:
                raise UsageError("Native OES workbook sheet layout changed; expected one sheet.")
            sheet = book.worksheets[0]
            # The observed OES producer declares an A1 dimension despite containing all data rows.
            sheet.reset_dimensions()
            values = list(sheet.iter_rows(values_only=True))
        finally:
            book.close()
    except UsageError:
        raise
    except Exception as exc:
        raise UsageError("Downloaded attachment is not a readable XLSX workbook.") from exc
    values = [row for row in values if any(cell is not None and str(cell).strip() for cell in row)]
    if not values:
        raise UsageError("Downloaded XLSX has no column headers.")
    headers = [str(cell).strip() if cell is not None else "" for cell in values[0]]
    rows = values[1:]
    if len(rows) != expected_count:
        raise UsageError(f"Native XLSX row count mismatch: expected {expected_count}, received {len(rows)}.")
    verification: dict[str, Any] = {"row_count": len(rows), "headers": headers, "row_count_verified": True,
                                    "order_numbers_verified": False}
    if query_rows is not None:
        if len(query_rows) != expected_count or any(not isinstance(row, dict) or not row.get("orderNumber") for row in query_rows):
            raise UsageError("Query rows do not contain a complete orderNumber contract.")
        order_headers = [i for i, header in enumerate(headers) if header in ("订单编号", "订单号", "订单编码")]
        if len(order_headers) != 1:
            raise UsageError("Native workbook's order-number column changed; refusing unverified query/mail reconciliation.")
        column = order_headers[0]
        api_orders = Counter(str(row["orderNumber"]).strip() for row in query_rows)
        file_orders = Counter(str(row[column]).strip() for row in rows)
        if api_orders != file_orders:
            raise UsageError("Native XLSX order numbers differ from the exact OES query result; delivery stopped.")
        verification["order_numbers_verified"] = True
        identity_headers = {"clazzBizNumber": "班级bizNumber", "userId": "userId", "price": "业绩金额"}
        if any(headers.count(header) != 1 for header in identity_headers.values()):
            raise UsageError("Native workbook's class/user/amount columns changed; query reconciliation stopped.")
        fields = ["orderNumber", *identity_headers]
        columns = [column, *(headers.index(header) for header in identity_headers.values())]
        if any(any(field not in row for field in fields) for row in query_rows):
            raise UsageError("Query rows lack class/user/amount reconciliation fields.")
        def canonical(values: list[Any]) -> tuple[Any, ...]:
            try:
                amount = Decimal(str(values[-1]).strip())
            except InvalidOperation as exc:
                raise UsageError("OES query or native workbook has a nonnumeric performance amount.") from exc
            if not amount.is_finite():
                raise UsageError("OES query or native workbook has a nonfinite performance amount.")
            # XLSX producers/readers may represent the same empty text cell as None or "".
            return (*("" if value is None else str(value).strip() for value in values[:-1]), amount)
        if Counter(canonical([row[field] for field in fields]) for row in query_rows) != Counter(canonical([row[i] for i in columns]) for row in rows):
            raise UsageError("Native XLSX class IDs, user IDs or amounts differ from the exact OES query result.")
        verification["reconciled_fields"] = fields
        verification["blank_text_cells_normalized"] = True
    return verification


def save_attachment(data: bytes, output_dir: Path, identity_hash: str) -> Path:
    output_dir = validate_output_dir(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    # Never trust an attachment's supplied filename as a filesystem path.
    path = output_dir / f"oes-native-{identity_hash[:20]}.xlsx"
    temporary = path.with_suffix(".part")
    temporary.write_bytes(data)
    temporary.replace(path)
    if path.read_bytes() != data:
        raise UsageError("Downloaded attachment failed local file readback.")
    return path


def merge_native_workbooks(paths: list[Path], query_rows: list[dict[str, Any]], output_dir: Path) -> dict[str, Any]:
    """Merge complete native batches while preserving literal IDs and source values."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    style_root = SKILLS_ROOT / "xlsx" / "scripts"
    if not style_root.is_dir():
        raise UsageError("The xlsx skill's formatting helpers are unavailable; native batches remain intact.")
    sys.path.insert(0, str(style_root))
    from style_apply import apply_header_style, apply_auto_fit_columns, apply_number_format, apply_border_grid

    if len(query_rows) > 1_048_575:
        raise UsageError("Merged data exceeds one Excel sheet's row capacity; verified native batches and JSON remain available.")
    headers = None
    samples: list[tuple[Any, ...]] = []
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style, apply openpyxl's default", category=UserWarning)
        first = load_workbook(paths[0], read_only=True, data_only=True)
        try:
            first.active.reset_dimensions()
            iterator = first.active.iter_rows(values_only=True)
            headers = tuple(next(iterator))
            for row in iterator:
                samples.append(row)
                if len(samples) >= 100:
                    break
        finally:
            first.close()
    format_book = Workbook()
    format_sheet = format_book.active
    format_sheet.append(headers)
    for row in samples:
        format_sheet.append(row)
    apply_header_style(format_sheet, row=1, max_col=len(headers))
    apply_auto_fit_columns(format_sheet, max_col=len(headers))
    apply_number_format(format_sheet, start_row=2, end_row=2)
    apply_border_grid(format_sheet, start_row=2, end_row=2, max_col=len(headers))
    book = Workbook(write_only=True)
    sheet = book.create_sheet("订单明细")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(query_rows) + 1}"
    for column in range(1, len(headers) + 1):
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = format_sheet.column_dimensions[letter].width
    text_columns = {i for i, header in enumerate(headers) if header in ("订单号", "订单编号", "班级bizNumber", "手机号", "userId")}
    # Register styles in the destination workbook once per column, then reuse
    # immutable style arrays instead of copying/hashing objects for every cell.
    def cached_style(source: Any) -> Any:
        template = WriteOnlyCell(sheet)
        template.font, template.fill = copy(source.font), copy(source.fill)
        template.border, template.alignment = copy(source.border), copy(source.alignment)
        template.number_format = source.number_format
        return copy(template._style)
    header_styles = [cached_style(format_sheet.cell(1, i + 1)) for i in range(len(headers))]
    body_styles = []
    body_font = Font(name="Arial", size=11)
    for i in range(len(headers)):
        template = format_sheet.cell(2, i + 1)
        template.font = body_font
        if i in text_columns:
            template.number_format = "@"
        body_styles.append(cached_style(template))
    def append(values: tuple[Any, ...], *, header: bool = False) -> None:
        cells = []
        for i, value in enumerate(values):
            if not header and i in text_columns and value is not None:
                value = str(value)
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
            cell._style = header_styles[i] if header else body_styles[i]
            cells.append(cell)
        sheet.append(cells)
    append(headers, header=True)
    try:
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Workbook contains no default style, apply openpyxl's default", category=UserWarning)
            for path in paths:
                native = load_workbook(path, read_only=True, data_only=True)
                try:
                    native.active.reset_dimensions()
                    iterator = native.active.iter_rows(values_only=True)
                    if tuple(next(iterator)) != headers:
                        raise UsageError("Native batch workbook headers differ; merged delivery stopped.")
                    for row in iterator:
                        if any(value is not None and str(value).strip() for value in row):
                            append(row)
                finally:
                    native.close()
        buffer = BytesIO()
        book.save(buffer)
    finally:
        format_book.close()
        book.close()
    data = buffer.getvalue()
    verification = verify_xlsx(data, len(query_rows), query_rows)
    path = validate_output_dir(output_dir) / "oes-achievement-combined.xlsx"
    temporary = path.with_suffix(".part")
    temporary.write_bytes(data)
    temporary.replace(path)
    if path.read_bytes() != data:
        raise UsageError("Merged workbook failed file readback.")
    return {"file": file_info(path), "verification": verification}
