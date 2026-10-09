"""Convert every native Excel field and reconcile the complete remote table."""

from collections import Counter
from datetime import datetime
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import warnings

from _shared.errors import UsageError
from .board import SHANGHAI

HEADERS = ("时间", "订单号", "班级bizNumber", "班级名称", "业绩归属类型", "用户姓名", "手机号", "userId",
           "业绩金额", "业绩归属人", "转介绍归属类型", "状态", "是否扣绩效")


def timestamp_ms(value) -> int:
    if isinstance(value, bool):
        raise UsageError("Invalid OES/Base datetime value.")
    if isinstance(value, (int, float)):
        return int(value)
    try:
        parsed = value if isinstance(value, datetime) else datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=SHANGHAI)
        return int(parsed.timestamp() * 1000)
    except (ValueError, TypeError, OverflowError) as exc:
        raise UsageError("OES/Base datetime is not a valid date, ISO time or millisecond timestamp.") from exc


def field_definitions() -> list[dict]:
    return [({"name": name, "type": "datetime", "style": {"format": "yyyy-MM-dd HH:mm"}} if name == "时间"
             else {"name": name, "type": "number", "style": {"type": "plain", "precision": 2}} if name == "业绩金额"
             else {"name": name, "type": "text", "style": {"type": "plain"}}) for name in HEADERS]


def validate_schema(fields: list[dict]) -> None:
    expected = {field["name"]: field["type"] for field in field_definitions()}
    actual = {field["name"]: field["type"] for field in fields}
    if actual != expected or len(fields) != len(expected):
        raise UsageError("Target schema must match all 13 Excel fields; initialize with base-setup before any deletion.")


def workbook_records(paths: list[Path], expected_count: int) -> list[dict]:
    from openpyxl import load_workbook
    records = []
    for path in paths:
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Workbook contains no default style, apply openpyxl's default")
            book = load_workbook(path, read_only=True, data_only=False)
        try:
            if len(book.worksheets) != 1:
                raise UsageError("Native Excel sheet layout changed.")
            sheet = book.active
            sheet.reset_dimensions()
            iterator = sheet.iter_rows()
            headers = tuple(cell.value for cell in next(iterator))
            if headers != HEADERS:
                raise UsageError("Native Excel headers differ from the 13-field Base contract.")
            for cells in iterator:
                if not any(cell.value is not None and str(cell.value).strip() for cell in cells):
                    continue
                if len(cells) != len(HEADERS) or any(cell.data_type == "f" for cell in cells):
                    raise UsageError("Native Excel has unexpected columns or formula cells.")
                row = {}
                for name, cell in zip(HEADERS, cells):
                    value = cell.value
                    if name == "时间":
                        value = timestamp_ms(value)
                    elif name == "业绩金额":
                        try:
                            number = Decimal(str(value))
                            value = float(number)
                            if not number.is_finite() or Decimal(str(value)) != number:
                                raise ValueError("lossy number")
                        except (InvalidOperation, ValueError, OverflowError) as exc:
                            raise UsageError("Excel performance amount cannot be represented faithfully in Base.") from exc
                    elif value is not None:
                        value = str(value)
                    row[name] = value
                records.append(row)
        finally:
            book.close()
    if len(records) != expected_count:
        raise UsageError("Combined native Excel count differs from the completed OES query.")
    return records


def canonical(row: dict) -> tuple:
    values = []
    for name in HEADERS:
        value = row.get(name)
        if name == "时间":
            value = None if value is None else timestamp_ms(value)
        elif name == "业绩金额":
            value = None if value is None else str(Decimal(str(value)).normalize())
        else:
            value = "" if value is None else str(value)
        values.append(value)
    return tuple(values)


def verify_records(expected: list[dict], actual: list[dict]) -> dict:
    left, right = Counter(map(canonical, expected)), Counter(map(canonical, actual))
    if left != right:
        raise UsageError(f"Full Base readback differs from all Excel fields: expected={len(expected)}, actual={len(actual)}, missing_rows={sum((left-right).values())}, extra_rows={sum((right-left).values())}.")
    digest = hashlib.sha256(json.dumps(sorted((json.dumps(key, ensure_ascii=True), count) for key, count in left.items()),
                                      separators=(",", ":")).encode()).hexdigest()
    return {"verified": True, "row_count": len(actual), "fields_verified": list(HEADERS),
            "cell_comparisons": len(actual) * len(HEADERS), "row_multiset_sha256": digest, "mismatch_count": 0}
