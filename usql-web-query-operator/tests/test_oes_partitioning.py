from datetime import date, datetime
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from _shared.errors import UsageError
from oes_achievement.board import TimeWindow, build_payload
from oes_achievement.cli import build_parser
from oes_achievement.partitioning import build_partition_plan, rows_in_window, run_partitioned_pipeline
from oes_achievement.workflow import resolve_job_options
from oes_achievement.workbook import atomic_json, file_info, merge_native_workbooks, verify_xlsx


def row(identity, timestamp, *, order="repeated-order"):
    return {"id": identity, "tradeTime": timestamp, "orderNumber": order, "clazzBizNumber": "class", "userId": "user", "price": "1.00"}


def workbook(path, rows):
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.append(["订单号", "班级bizNumber", "userId", "业绩金额"])
    for value in rows:
        sheet.append([value["orderNumber"], value["clazzBizNumber"], value["userId"], value["price"]])
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    book.close()


class TimePartitionTests(unittest.TestCase):
    def test_adjacent_ranges_cover_midnight_without_gap_or_overlap(self):
        window = TimeWindow.from_dates(date(2026, 10, 1), date(2026, 10, 2))
        midnight = TimeWindow.from_dates(date(2026, 10, 2), date(2026, 10, 2)).start_ms
        rows = [row(1, window.start_ms), row(2, midnight - 1), row(3, midnight), row(4, window.end_ms)]
        plan = build_partition_plan(rows, window, limit=2)
        self.assertEqual(len(plan), 2)
        self.assertEqual(plan[0]["end_ms"] + 1, plan[1]["start_ms"])
        self.assertEqual(plan[0]["end_ms"], midnight - 1)
        self.assertEqual(plan[1]["start_ms"], midnight)
        collected = [item["id"] for part in plan for item in rows_in_window(rows, TimeWindow(part["start_ms"], part["end_ms"]))]
        self.assertCountEqual(collected, [1, 2, 3, 4])
        self.assertEqual(len(set(collected)), len(collected))

    def test_single_busy_day_splits_below_one_day(self):
        window = TimeWindow.from_dates(date(2026, 10, 1), date(2026, 10, 1))
        span = window.end_ms - window.start_ms
        rows = [row(i, window.start_ms + span * i // 8) for i in range(9)]
        plan = build_partition_plan(rows, window, limit=2)
        self.assertGreater(len(plan), 1)
        self.assertEqual(plan[0]["start_ms"], window.start_ms)
        self.assertEqual(plan[-1]["end_ms"], window.end_ms)
        self.assertTrue(all(part["row_count"] <= 2 for part in plan))
        self.assertEqual(sum(part["row_count"] for part in plan), 9)

    def test_identical_order_numbers_are_preserved_as_distinct_rows(self):
        window = TimeWindow(0, 100)
        plan = build_partition_plan([row(1, 0), row(2, 100)], window, limit=1)
        self.assertEqual(sum(part["row_count"] for part in plan), 2)

    def test_unsplittable_same_millisecond_is_blocked(self):
        with self.assertRaisesRegex(UsageError, "share one millisecond"):
            build_partition_plan([row(i, 50) for i in range(3)], TimeWindow(0, 100), limit=2)

    def test_out_of_range_or_duplicate_source_ids_are_blocked(self):
        for rows in ([row(1, -1)], [row(1, 0), row(1, 50)]):
            with self.subTest(rows=rows), self.assertRaises(UsageError):
                build_partition_plan(rows, TimeWindow(0, 100), limit=1)

    def test_intraday_payload_preserves_exact_millisecond_bounds(self):
        window = TimeWindow(1790784000001, 1790784000002)
        payload = build_payload(*window.dates, page_num=1, page_size=1, window=window)
        self.assertEqual(payload["beginTime"], str(window.start_ms))
        self.assertEqual(payload["endTime"], str(window.end_ms))


class MergedWorkbookTests(unittest.TestCase):
    def test_merge_keeps_blank_class_ids_empty_without_false_reconciliation_failure(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            rows = [dict(row(1, 0), clazzBizNumber="")]
            source = root / "blank.xlsx"
            workbook(source, rows)
            result = merge_native_workbooks([source], rows, root)
            self.assertTrue(result["verification"]["blank_text_cells_normalized"])
            self.assertEqual(result["verification"]["row_count"], 1)

    def test_merge_retains_duplicates_text_ids_and_has_verified_format(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            rows = [row(1, 0, order="001234567890123456789"), row(2, 100, order="001234567890123456789")]
            paths = [root / "one.xlsx", root / "two.xlsx"]
            for path, value in zip(paths, rows):
                workbook(path, [value])
            result = merge_native_workbooks(paths, rows, root)
            self.assertEqual(result["verification"]["row_count"], 2)
            self.assertTrue(result["verification"]["order_numbers_verified"])
            book = load_workbook(result["file"]["path"])
            try:
                self.assertEqual(book.active.freeze_panes, "A2")
                self.assertEqual(book.active["A2"].value, "001234567890123456789")
                self.assertEqual(book.active["A3"].value, "001234567890123456789")
                self.assertEqual(book.active["A2"].data_type, "s")
                self.assertEqual(book.active["A1"].fill.patternType, "solid")
                self.assertEqual(book.active["A2"].border.bottom.style, "thin")
                self.assertEqual(book.active["A2"].font.name, "Arial")
                self.assertEqual(book.active["A2"].font.sz, 11)
                self.assertEqual(book.active["A2"].number_format, "@")
                self.assertEqual(book.active["A2"]._style, book.active["A3"]._style)
            finally:
                book.close()


class BatchWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.args = build_parser().parse_args(["export-and-download", "--start-date", "2026-10-01", "--end-date", "2026-10-02", "--output-dir", str(self.root / "output"), "--run-key", "large-run"])
        self.start, self.end = resolve_job_options(self.args)
        self.window = TimeWindow.from_dates(self.start, self.end)
        second = TimeWindow.from_dates(self.end, self.end).start_ms
        self.rows = [row(i, self.window.start_ms if i % 2 else second) for i in range(1, 10003)]
        self.exported = []
        self.fail_second = False
        self.calls = {
            "OES_RUNTIME_DIR": self.root / "runtime",
            "open_session": Mock(return_value=Mock()),
            "inspect_range": Mock(side_effect=self.inspect),
            "query_all": Mock(return_value=(self.rows, len(self.rows), [])),
            "run_pipeline": Mock(side_effect=self.native_batch),
        }
        for name, value in self.calls.items():
            active = patch("oes_achievement.partitioning." + name, value)
            active.start()
            self.addCleanup(active.stop)

    def inspect(self, request, start, end, timeout, *, window=None):
        return {"row_count": len(rows_in_window(self.rows, window)) if window else len(self.rows)}

    def native_batch(self, playwright, args, start, end):
        import hashlib
        directory = args.output_dir / ("oes-run-" + hashlib.sha256(args.run_key.encode()).hexdigest()[:24])
        receipt_path = directory / "receipt.json"
        if receipt_path.exists():
            import json
            return json.loads(receipt_path.read_text(encoding="utf-8"))
        if self.fail_second and self.exported:
            raise UsageError("Second batch mail deadline")
        self.exported.append(args.run_key)
        path = directory / "native.xlsx"
        workbook(path, args.prepared_rows)
        result = {"status": "completed", "receipt_path": str(receipt_path), "mail_download": {
            "file": file_info(path), "verification": verify_xlsx(path.read_bytes(), len(args.prepared_rows), args.prepared_rows)}}
        atomic_json(receipt_path, result)
        return result

    def test_more_than_10000_rows_export_in_batches_and_merge_exactly(self):
        result = run_partitioned_pipeline(None, self.args, self.start, self.end)
        self.assertEqual(result["status"], "completed")
        self.assertEqual(result["row_count"], 10002)
        self.assertEqual(result["batch_count"], 2)
        self.assertEqual(result["merged_workbook"]["verification"]["row_count"], 10002)
        self.assertEqual(len(self.exported), 2)
        self.assertEqual(result["export_dispatch"], "serial_for_mail_disambiguation")
        repeated = run_partitioned_pipeline(None, self.args, self.start, self.end)
        self.assertTrue(repeated["reused_completed_run"])
        self.assertEqual(len(self.exported), 2)

    def test_distinct_count_exports_are_submitted_before_receiving_any_mail(self):
        self.rows.append(row(10003, self.window.start_ms))
        self.calls["query_all"].return_value = (self.rows, len(self.rows), [])
        result = run_partitioned_pipeline(None, self.args, self.start, self.end)
        modes = [getattr(call.args[1], "submit_only", False) for call in self.calls["run_pipeline"].call_args_list]
        self.assertEqual(modes, [True, True, False, False])
        self.assertEqual(len(self.exported), 2)
        self.assertEqual(result["export_dispatch"], "submit_all_then_receive")
        self.assertEqual(result["merged_workbook"]["verification"]["row_count"], len(self.rows))

    def test_partial_batch_resume_does_not_resubmit_completed_batch(self):
        self.fail_second = True
        with self.assertRaisesRegex(UsageError, "same --run-key"):
            run_partitioned_pipeline(None, self.args, self.start, self.end)
        self.fail_second = False
        result = run_partitioned_pipeline(None, self.args, self.start, self.end)
        self.assertEqual(result["status"], "completed")
        self.assertEqual(len(self.exported), 2)
        self.assertEqual(self.calls["query_all"].call_count, 1)

    def test_api_partition_count_mismatch_stops_before_any_native_batch(self):
        self.calls["inspect_range"].side_effect = lambda request, start, end, timeout, window=None: {"row_count": 1 if window else len(self.rows)}
        with self.assertRaisesRegex(UsageError, "do not match live API counts"):
            run_partitioned_pipeline(None, self.args, self.start, self.end)
        self.assertEqual(self.exported, [])
        self.calls["run_pipeline"].assert_not_called()


if __name__ == "__main__":
    unittest.main()
